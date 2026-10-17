from io import BytesIO
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
import unicodedata
import warnings

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
import xlrd

from database.reports import get_orders


def _parse_order_amount(value: object, row_number: int, worksheet_name: str) -> Decimal:
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))

    if isinstance(value, str):
        normalized_value = value.strip()
        if "," in normalized_value:
            normalized_value = normalized_value.replace(".", "").replace(",", ".")
        try:
            return Decimal(normalized_value)
        except InvalidOperation as error:
            raise ValueError(
                f"Valor total inválido na linha {row_number} da aba "
                f"'{worksheet_name}'."
            ) from error

    raise ValueError(
        f"Valor total inválido na linha {row_number} da aba '{worksheet_name}'."
    )


def summarize_orders_by_operator(
    file_content: bytes,
) -> tuple[
    dict[date, dict[str, dict[str, int]]],
    dict[str, dict[str, int]],
    str,
]:
    if file_content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        return _summarize_sitcarga_orders(file_content)

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="^Workbook contains no default style, apply openpyxl's default$",
            category=UserWarning,
            module=r"openpyxl\.styles\.stylesheet",
        )
        workbook = load_workbook(
            BytesIO(file_content),
            read_only=True,
            data_only=True,
        )

    try:
        if not workbook.worksheets:
            raise ValueError("A planilha não contém abas.")

        worksheet = workbook.worksheets[0]
        daily_counts: defaultdict[date, defaultdict[str, Counter[str]]] = defaultdict(
            lambda: defaultdict(Counter)
        )
        totals: defaultdict[str, Counter[str]] = defaultdict(Counter)

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=9,
                values_only=True,
            ),
            start=9,
        ):
            order_date = row[2] if len(row) > 2 else None
            operator = row[6] if len(row) > 6 else None
            order_amount = row[15] if len(row) > 15 else None

            if order_date is None and operator is None:
                continue
            if (
                order_date is None
                or not isinstance(operator, str)
                or not operator.strip()
            ):
                raise ValueError(
                    f"Registro incompleto na linha {row_number} da aba "
                    f"'{worksheet.title}'."
                )

            if isinstance(order_date, datetime):
                parsed_date = order_date.date()
            elif isinstance(order_date, date):
                parsed_date = order_date
            elif isinstance(order_date, str):
                try:
                    parsed_date = datetime.strptime(
                        order_date.strip(),
                        "%d/%m/%Y",
                    ).date()
                except ValueError as error:
                    raise ValueError(
                        f"Data inválida na linha {row_number} da aba "
                        f"'{worksheet.title}'."
                    ) from error
            else:
                raise ValueError(
                    f"Data inválida na linha {row_number} da aba "
                    f"'{worksheet.title}'."
                )

            operator_name = operator.strip()
            amount = _parse_order_amount(order_amount, row_number, worksheet.title)
            category = "free" if amount == 0 else "paid"
            daily_counts[parsed_date][operator_name][category] += 1
            totals[operator_name][category] += 1

        return (
            {
                day: {
                    operator: dict(counts)
                    for operator, counts in operator_counts.items()
                }
                for day, operator_counts in sorted(daily_counts.items())
            },
            {operator: dict(counts) for operator, counts in totals.items()},
            "pf",
        )
    finally:
        workbook.close()


def _normalize_header(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    return (
        "".join(
            character
            for character in normalized
            if not unicodedata.combining(character)
        )
        .strip()
        .casefold()
    )


def _summarize_sitcarga_orders(
    file_content: bytes,
) -> tuple[
    dict[date, dict[str, dict[str, int]]],
    dict[str, dict[str, int]],
    str,
]:
    try:
        workbook = xlrd.open_workbook(file_contents=file_content, on_demand=True)
    except (xlrd.XLRDError, OSError) as error:
        raise ValueError("Não foi possível abrir o arquivo XLS.") from error

    try:
        if not workbook.nsheets:
            raise ValueError("O arquivo não contém abas.")

        worksheet = workbook.sheet_by_index(0)
        header_row = None
        column_indexes = {}
        required_headers = {
            "tipo servico": "service",
            "data": "date",
            "servico (r$)": "amount",
            "nome operador": "operator",
        }

        for row_index in range(min(worksheet.nrows, 20)):
            headers = {
                _normalize_header(worksheet.cell_value(row_index, column_index)): (
                    column_index
                )
                for column_index in range(worksheet.ncols)
            }
            if all(header in headers for header in required_headers):
                header_row = row_index
                column_indexes = {
                    key: headers[header] for header, key in required_headers.items()
                }
                break

        if header_row is None:
            raise ValueError(
                "O arquivo XLS não corresponde ao formato de relatório esperado."
            )

        daily_counts: defaultdict[date, defaultdict[str, Counter[str]]] = defaultdict(
            lambda: defaultdict(Counter)
        )
        totals: defaultdict[str, Counter[str]] = defaultdict(Counter)

        for row_index in range(header_row + 1, worksheet.nrows):
            row = worksheet.row(row_index)
            order_date_cell = row[column_indexes["date"]]
            operator_value = worksheet.cell_value(
                row_index,
                column_indexes["operator"],
            )
            service_value = worksheet.cell_value(
                row_index,
                column_indexes["service"],
            )
            amount_cell = row[column_indexes["amount"]]

            if (
                order_date_cell.ctype == xlrd.XL_CELL_EMPTY
                and not operator_value
                and not service_value
            ):
                continue

            row_number = row_index + 1
            amount = _parse_order_amount(
                amount_cell.value,
                row_number,
                worksheet.name,
            )
            if amount == 0:
                continue

            if (
                order_date_cell.ctype == xlrd.XL_CELL_EMPTY
                or not isinstance(operator_value, str)
                or not operator_value.strip()
            ):
                raise ValueError(
                    f"Registro incompleto na linha {row_number} da aba "
                    f"'{worksheet.name}'."
                )

            try:
                if order_date_cell.ctype == xlrd.XL_CELL_DATE:
                    parsed_date = xlrd.xldate_as_datetime(
                        order_date_cell.value,
                        workbook.datemode,
                    ).date()
                elif isinstance(order_date_cell.value, str):
                    raw_date = order_date_cell.value.strip()
                    try:
                        parsed_date = datetime.fromisoformat(raw_date).date()
                    except ValueError:
                        parsed_date = datetime.strptime(
                            raw_date,
                            "%d/%m/%Y",
                        ).date()
                else:
                    raise ValueError
            except (ValueError, xlrd.XLDateError) as error:
                raise ValueError(
                    f"Data inválida na linha {row_number} da aba "
                    f"'{worksheet.name}'."
                ) from error

            operator_name = operator_value.strip()
            daily_counts[parsed_date][operator_name]["paid"] += 1
            totals[operator_name]["paid"] += 1

        return (
            {
                day: {
                    operator: dict(counts)
                    for operator, counts in operator_counts.items()
                }
                for day, operator_counts in sorted(daily_counts.items())
            },
            {operator: dict(counts) for operator, counts in totals.items()},
            "pj",
        )
    finally:
        workbook.release_resources()


def create_operator_analysis_workbook(
    daily_counts: dict[date, dict[str, dict[str, int]]],
    totals: dict[str, dict[str, int]],
    discord_paid_operator_counts: dict[date, dict[str, int]],
    discord_free_operator_counts: dict[date, dict[str, int]],
) -> bytes:
    workbook = Workbook()
    daily_sheet = workbook.active
    daily_sheet.title = "Totais diários"
    operator_daily_sheet = workbook.create_sheet("Por operador")
    operator_totals_sheet = workbook.create_sheet("Totais por operador")

    first_day = min(daily_counts)
    last_day = max(daily_counts)

    operator_labels = {}
    normalized_daily_counts = {}
    normalized_totals = {}
    normalized_discord_paid_counts = {}
    normalized_discord_free_counts = {}

    for day, operators in daily_counts.items():
        normalized_operators = normalized_daily_counts.setdefault(day, {})
        for name, counts in operators.items():
            key = _normalize_header(name)
            operator_labels.setdefault(key, name)
            normalized_counts = normalized_operators.setdefault(key, {})
            for category, count in counts.items():
                normalized_counts[category] = normalized_counts.get(category, 0) + count

    for name, counts in totals.items():
        key = _normalize_header(name)
        operator_labels.setdefault(key, name)
        normalized_counts = normalized_totals.setdefault(key, {})
        for category, count in counts.items():
            normalized_counts[category] = normalized_counts.get(category, 0) + count

    def normalize_discord_counts(counts_by_day):
        normalized_counts_by_day = {}
        for day, operators in counts_by_day.items():
            normalized_operators = normalized_counts_by_day.setdefault(day, {})
            for name, count in operators.items():
                key = _normalize_header(name)
                operator_labels.setdefault(key, name)
                normalized_operators[key] = normalized_operators.get(key, 0) + count
        return normalized_counts_by_day

    normalized_discord_paid_counts = normalize_discord_counts(
        discord_paid_operator_counts
    )
    normalized_discord_free_counts = normalize_discord_counts(
        discord_free_operator_counts
    )

    daily_headers = [
        "Data",
        "Com valor (planilha)",
        "Gratuitos (planilha)",
        "Total planilha",
        "Pagos registrados no Discord",
        "Gratuitos registrados no Discord",
        "Total registrados no Discord",
    ]
    daily_sheet.append(daily_headers)
    paid_total = 0
    free_total = 0
    discord_paid_total = 0
    discord_free_total = 0
    day = first_day
    while day <= last_day:
        counts = daily_counts.get(day, {})
        paid = sum(
            operator_counts.get("paid", 0) for operator_counts in counts.values()
        )
        free = sum(
            operator_counts.get("free", 0) for operator_counts in counts.values()
        )
        discord_paid = sum(normalized_discord_paid_counts.get(day, {}).values())
        discord_free = sum(normalized_discord_free_counts.get(day, {}).values())
        daily_sheet.append(
            [
                day,
                paid,
                free,
                paid + free,
                discord_paid,
                discord_free,
                discord_paid + discord_free,
            ]
        )
        paid_total += paid
        free_total += free
        discord_paid_total += discord_paid
        discord_free_total += discord_free
        day += timedelta(days=1)

    daily_sheet.append(
        [
            "TOTAL DO PERÍODO",
            paid_total,
            free_total,
            paid_total + free_total,
            discord_paid_total,
            discord_free_total,
            discord_paid_total + discord_free_total,
        ]
    )

    operators = set(operator_labels)
    operator_order = sorted(
        operators,
        key=lambda operator: (
            -(
                normalized_totals.get(operator, {}).get("paid", 0)
                + normalized_totals.get(operator, {}).get("free", 0)
            ),
            operator.casefold(),
        ),
    )

    operator_daily_sheet.append(
        [
            "Operador",
            "Data",
            "Com valor (planilha)",
            "Gratuitos (planilha)",
            "Total (planilha)",
            "Pagos (Discord)",
            "Gratuitos (Discord)",
            "Total (Discord)",
        ]
    )
    for operator in operator_order:
        day = first_day
        while day <= last_day:
            counts = normalized_daily_counts.get(day, {}).get(operator, {})
            paid = counts.get("paid", 0)
            free = counts.get("free", 0)
            discord_paid = normalized_discord_paid_counts.get(day, {}).get(operator, 0)
            discord_free = normalized_discord_free_counts.get(day, {}).get(operator, 0)
            operator_daily_sheet.append(
                [
                    operator_labels[operator],
                    day,
                    paid,
                    free,
                    paid + free,
                    discord_paid,
                    discord_free,
                    discord_paid + discord_free,
                ]
            )
            day += timedelta(days=1)

    operator_totals_sheet.append(
        [
            "Operador",
            "Com valor (planilha)",
            "Gratuitos (planilha)",
            "Total (planilha)",
            "Pagos (Discord)",
            "Gratuitos (Discord)",
            "Total (Discord)",
        ]
    )
    for operator in operator_order:
        counts = normalized_totals.get(operator, {})
        paid = counts.get("paid", 0)
        free = counts.get("free", 0)
        discord_paid = sum(
            counts_by_day.get(operator, 0)
            for counts_by_day in normalized_discord_paid_counts.values()
        )
        discord_free = sum(
            counts_by_day.get(operator, 0)
            for counts_by_day in normalized_discord_free_counts.values()
        )
        operator_totals_sheet.append(
            [
                operator_labels[operator],
                paid,
                free,
                paid + free,
                discord_paid,
                discord_free,
                discord_paid + discord_free,
            ]
        )

    _style_analysis_sheet(
        daily_sheet,
        [16, 24, 24, 18, 28, 31, 30],
        total_row=daily_sheet.max_row,
    )
    _style_analysis_sheet(
        operator_daily_sheet,
        [40, 16, 24, 24, 18, 20, 23, 18],
    )
    _style_analysis_sheet(
        operator_totals_sheet,
        [40, 24, 24, 18, 20, 23, 18],
    )

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _style_analysis_sheet(
    worksheet: Worksheet,
    column_widths: list[int],
    total_row: int | None = None,
) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    total_fill = PatternFill("solid", fgColor="D9EAF7")
    alternate_fill = PatternFill("solid", fgColor="F3F6FA")
    bottom_border = Border(bottom=Side(style="thin", color="D9E2F3"))

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 32

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    for row_number in range(2, worksheet.max_row + 1):
        is_total = row_number == total_row
        fill = (
            total_fill
            if is_total
            else (alternate_fill if row_number % 2 == 0 else None)
        )
        for cell in worksheet[row_number]:
            if fill is not None:
                cell.fill = fill
            cell.border = bottom_border
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if is_total:
                cell.font = Font(bold=True)
        if isinstance(worksheet.cell(row_number, 1).value, date):
            worksheet.cell(row_number, 1).number_format = "dd/mm/yyyy"
        if worksheet.title == "Por operador" and isinstance(
            worksheet.cell(row_number, 2).value,
            date,
        ):
            worksheet.cell(row_number, 2).number_format = "dd/mm/yyyy"

    for column_number, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[get_column_letter(column_number)].width = width

    if total_row is None:
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(column_widths))}{worksheet.max_row}"
        )
    elif total_row > 2:
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(column_widths))}{total_row - 1}"
        )


def generate_daily_excel(
    *,
    guild_id: str,
    start_date: datetime,
    end_date: datetime,
) -> str:

    workbook = Workbook()

    pf_sheet = workbook.active
    pf_sheet.title = "PF"

    pj_sheet = workbook.create_sheet(
        "PJ",
    )

    pf_sheet.append(
        [
            "Data",
            "Cliente",
            "CPF",
            "Pedidos",
            "Qtd. Taxas",
            "Cursos",
            "Valor",
            "Operador",
            "Observações",
        ]
    )

    pj_sheet.append(
        [
            "Data",
            "Cliente",
            "CNPJ",
            "Pedidos",
            "Cadastros/Reativações",
            "Alterações/Exclusões",
            "Cursos",
            "Valor",
            "Operador",
            "Observações",
        ]
    )

    orders = get_orders(
        guild_id=guild_id,
        start_date=start_date,
        end_date=end_date,
    )

    for order in orders:

        if order.category == "pf":

            pf_sheet.append(
                [
                    order.finished_at.strftime("%d/%m/%Y %H:%M"),
                    order.client,
                    order.document,
                    order.order,
                    order.pf_amount,
                    order.course_amount,
                    order.dispatcher_value,
                    order.operator_name,
                    order.observations or "",
                ]
            )

        elif order.category == "pj":

            pj_sheet.append(
                [
                    order.finished_at.strftime("%d/%m/%Y %H:%M"),
                    order.client,
                    order.document,
                    order.order,
                    order.pj_amount_cad_or_reval,
                    order.pj_amount_alt_or_rem,
                    order.course_amount,
                    order.dispatcher_value,
                    order.operator_name,
                    order.observations or "",
                ]
            )

    filename = (
        Path(__file__).parent
        / f"fechamento-{guild_id}-{end_date.strftime('%Y-%m-%d')}.xlsx"
    )

    workbook.save(
        filename,
    )

    workbook.close()

    return str(
        filename,
    )
